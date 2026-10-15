"""SPARK lot 4 — convertit le fichier « MAEVA Vendeurs - affaires.xlsx » en JSON pour le chargement initial du référentiel.

  python -I scripts/spark-vendeurs-maeva.py "<MAEVA…xlsx>" "<sortie.json>"
puis, sur le poste :  npx ts-node --transpile-only src/rcu/worker.ts --vendeurs "<sortie.json>"   (dans backend/)
et sur le VPS      :  sudo docker compose exec rcu-worker node dist/rcu/worker.js --vendeurs /archive/manuel/<sortie.json>

Source : feuille « Infos » (PLAQUE, SITE, RESPONSABLE SITE, SERVICE, VENDEUR, TYPE, NOTES) + feuille « ipn » (numéros D…).
Le fichier ne contient que des données de salariés : le JSON produit ne se commite pas (il reste hors du dépôt).
Après le chargement, c'est l'écran SPARK › Équipes qui fait foi, plus le fichier.
"""
import json, re, sys, unicodedata, warnings
import openpyxl

warnings.filterwarnings('ignore')
SITES = {'CLERMONT': 'Clermont', 'MOZAC': 'Mozac', 'MASSAGETTES': 'Massagettes', 'USSEL': 'Ussel', 'VICHY': 'Vichy', 'MOULINS': 'Moulins',
         'THIERS': 'Thiers', 'ISSOIRE': 'Issoire', 'MENDE': 'Mende', 'LE PUY': 'Le Puy-en-Velay', 'ALBI': 'Albi', 'RODEZ': 'Rodez',
         'MILLAU': 'Millau', 'AURILLAC': 'Aurillac', 'FIGEAC': 'Figeac', 'GAILLAC': 'Gaillac', 'VDR': 'Villefranche', 'CARMAUX': 'Carmaux',
         'LAVAUR': 'Lavaur'}
# Reprises de clients écrites en clair dans les notes → règles « ancien vendeur ».
REPRISES = [(r'remplace\s+fabrice\s+ranson', 'RANSON'), (r'reprend secteur j\.?\s*zarlenga', 'ZARLENGA'), (r'affecter robin trouche', 'TROUCHE')]
SORTIE_PARTI = '2026-10-14'   # « Parti pas remplacé » : date exacte inconnue, ses clients sont réattribués au premier calcul

plain = lambda s: unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().upper().strip()
low = lambda s: plain(s).lower()


def nom_prenom(v):
    w = str(v).split()
    nom = []
    while w and w[0] == w[0].upper() and any(ch.isalpha() for ch in w[0]):
        nom.append(w.pop(0))
    if not w:  # tout en majuscules (« ELOI MARCHAND ») : premier mot = nom
        return nom[0], ' '.join(x.capitalize() for x in nom[1:])
    return ' '.join(nom), ' '.join(w)


def lire(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
    lignes = [r for r in wb['Infos'].iter_rows(values_only=True) if any(c is not None for c in r)][1:]
    # Feuille « ipn » : un numéro D… est collé à un nom, à droite en 1re colonne, à gauche sinon.
    ipn = {}
    for r in wb['ipn'].iter_rows(values_only=True):
        c = [None if x is None else str(x).strip() for x in r]
        for i, x in enumerate(c):
            if x and re.fullmatch(r'[Dd]\d{5,6}', x):
                j = i + 1 if i == 0 else i - 1
                if 0 <= j < len(c) and c[j] and not re.fullmatch(r'[Dd]\d+', c[j]):
                    ipn[plain(re.sub(r'\(.*?\)', '', c[j]))] = x.upper()
    vendeurs, encadrement, regles, alertes = [], {}, [], []
    for plaque, site_m, resp, service, vendeur, typ, notes in [tuple(list(r) + [None] * 7)[:7] for r in lignes]:
        site = SITES.get(plain(site_m or ''))
        if not site:
            alertes.append(f'Site inconnu : {site_m}'); continue
        n = low(notes or '')
        if resp and service == 'VN':
            encadrement[(site, 'chef_de_site')] = re.sub(r'\s*\(.*?\)', '', str(resp)).strip()
        if resp and service == 'VO' and (site, 'chef_de_vente_vo') not in encadrement and 'cvvo' in n:
            encadrement[(site, 'chef_de_vente_vo')] = str(resp).strip()
        if not vendeur:
            continue
        if 'responsable vo' in n:
            encadrement[(site, 'chef_de_vente_vo')] = str(vendeur).strip(); continue
        nom, prenom = nom_prenom(vendeur)
        tels = re.findall(r'0\d(?: \d\d){4}', str(notes or ''))
        sans_tel = re.sub(r'0\d(?: \d\d){4}', ' ', n)
        zones = sorted(set(re.findall(r'(?<![\d\w])(\d{5}|\d{2})(?![\d\w])', sans_tel)))
        t = plain(typ or '')
        if 'alternant' in n: fonction = 'alternant'
        elif 'magasin' in n: fonction = 'magasin'
        elif 'secteur' in n or 'yssingeaux' in n: fonction = 'secteur'
        elif 'societe' in n or t == 'SO': fonction = 'societes'
        else: fonction = 'hall'
        marques = [] if service == 'VO' or t == 'VN-VO' else (['RENAULT', 'DACIA'] if t in ('RENAULT+DACIA', 'SO') else [t] if t in ('RENAULT', 'DACIA') else [])
        ip = ipn.get(plain(f'{nom} {prenom}')) or ipn.get(plain(f'{prenom} {nom}'))
        m = re.search(r'\b(d\d{6})\b', n)
        v = {'nom': nom, 'prenom': prenom, 'site': site, 'service': service, 'fonction': fonction, 'marques': marques,
             'recoit': fonction not in ('alternant', 'magasin'), 'ipn': ip or (m.group(1).upper() if m else None),
             'telephone': tels[0] if tels else None, 'zones': zones if fonction in ('secteur', 'magasin') else [],
             'note': str(notes).strip() if notes else None, 'sortie': SORTIE_PARTI if 'parti pas remplace' in n else None}
        vendeurs.append(v)
        for motif, ancien in REPRISES:
            if re.search(motif, n):
                regles.append({'site': site, 'type': 'ancien_vendeur', 'valeurs': [ancien], 'vendeur': f'{nom} {prenom}',
                               'note': f'Clients de {ancien.capitalize()} → {prenom} {nom} (fichier MAEVA)'})
    return {'vendeurs': vendeurs, 'encadrement': [{'site': s, 'role': r, 'nom': n} for (s, r), n in encadrement.items()],
            'regles': regles, 'alertes': alertes}


if __name__ == '__main__':
    out = lire(sys.argv[1])
    with open(sys.argv[2], 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    v = out['vendeurs']
    print(f"{len(v)} vendeurs ({sum(x['ipn'] is not None for x in v)} avec IPN, {sum(bool(x['zones']) for x in v)} avec secteur), "
          f"{len(out['encadrement'])} encadrants, {len(out['regles'])} règles, alertes : {out['alertes']}")
